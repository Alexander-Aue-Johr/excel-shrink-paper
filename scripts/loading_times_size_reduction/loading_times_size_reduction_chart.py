#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
import time
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional

import certifi
import numpy as np
import pandas as pd
import psutil
from openpyxl import load_workbook

import matplotlib.pyplot as plt
import matplotlib.transforms as mtransforms
from matplotlib.gridspec import GridSpec, GridSpecFromSubplotSpec
from collections import OrderedDict

plt.switch_backend("agg")

logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)
log = logging.getLogger(__name__)

Seconds = float
Bytes = int

# ----------------------------
# Paths / constants
# ----------------------------
SCRIPT_DIR = Path(__file__).resolve().parent

CHART_OUTPUT_PATH: Path = (
    SCRIPT_DIR / "charts" / "time_and_filesize_comparison_by_file.pdf"
)

EXCEL_SHRINK_SCRIPT: Path = (
    SCRIPT_DIR.parent / "excel_shrink_rust" / "run_excel_shrink_rust.py"
)
PYTHON_EXCEL_SHRINK_SCRIPT: Path = SCRIPT_DIR.parent / "excel_shrink" / "excel_shrink.py"
OLD_ANALYSER_SCRIPT: Path = (
    SCRIPT_DIR.parent.parent / "excel_shrink_analyzer" / "excel_shrink_analyzer.py"
)

EXCEL_SHRINK_SINGLE_CORE_ARGS = ["--disable-multiprocessing"]
EXCEL_SHRINK_NO_SPLIT_ARGS = ["--disable-sheetdata-splitting"]
EXCEL_SHRINK_RUST_DEFAULT_ARGS = [
    "--preprocess-hidden-rows",
    "--preprocess-empty-cells",
    "--remove-unreferenced-whitespace-cells",
    "--remove-unreferenced-hidden-content-cells",
]
DEFAULT_SHRINK_LIBRARY = "Excel Shrink Rust"


# ----------------------------
# Small utilities
# ----------------------------
def sanitize_filename(name: str) -> str:
    name = name.strip().rstrip(". ")
    return re.sub(r'[\\/:*?"<>|]', "_", name)


def maybe_size(p: Path) -> Optional[int]:
    try:
        return p.stat().st_size if p.exists() else None
    except OSError:
        return None


def measure_one_main(args: argparse.Namespace) -> None:
    library_name: str = args.library
    file_path = Path(args.file).resolve()
    on_shrunk: bool = bool(args.on_shrunk)

    bench = BENCHMARKS.get(library_name)
    if bench is None:
        log.error("Library %r not recognized.", library_name)
        print(json.dumps({"error": "Unknown library"}))
        sys.exit(1)

    out_path = (
        file_path.parent.parent
        / "library_output"
        / ("shrunk" if on_shrunk else "original")
        / sanitize_filename(library_name)
        / file_path.name
    ).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)

    open_t, save_t = bench(file_path, out_path)

    result = {
        "library": library_name,
        "file": file_path.name,
        "Original Open Time (s)": open_t,
        "Original Save Time (s)": save_t,
        "On Shrunk": on_shrunk,
    }
    print(json.dumps(result))


def measure_open_save(
    name: str, opener: Callable[[], None], saver: Callable[[], None]
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    try:
        t0 = time.perf_counter()
        opener()
        t_open = time.perf_counter() - t0

        t0 = time.perf_counter()
        saver()
        t_save = time.perf_counter() - t0

        return t_open, t_save
    except Exception:
        log.exception("%s benchmark error", name)
        return None, None


# ----------------------------
# Benchmarks (measure-one)
# ----------------------------
def benchmark_openpyxl_default(
    file_path: Path, out_path: Path
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    wb = None

    def _open() -> None:
        nonlocal wb
        wb = load_workbook(file_path)

    def _save() -> None:
        nonlocal wb
        assert wb is not None
        wb.save(out_path)
        log.info("openpyxl saved to %s", out_path)

    return measure_open_save("openpyxl", _open, _save)


def benchmark_pandas(
    file_path: Path, out_path: Path
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    dfs: dict[str, pd.DataFrame] = {}

    def _open() -> None:
        nonlocal dfs
        dfs = pd.read_excel(file_path, sheet_name=None)

    def _save() -> None:
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            for sheet, df in dfs.items():
                df.to_excel(writer, sheet_name=sheet, index=False)
        log.info("Pandas saved to %s", out_path)

    return measure_open_save("pandas", _open, _save)


def benchmark_excel_com(
    file_path: Path, out_path: Path
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    """
    Benchmark via win32com (Windows only).
    First runs excel_shrink.py --only-clean-workbook on the original,
    then opens the cleaned workbook with Excel COM and saves it.
    """
    try:
        import win32com.client  # type: ignore
    except ImportError:
        log.warning("win32com not available - skipping Excel COM benchmark.")
        return None, None

    with tempfile.TemporaryDirectory() as shrink_dir_str:
        shrink_dir = Path(shrink_dir_str)

        excel = None
        wb = None

        def _open() -> None:
            nonlocal excel, wb
            log.info("Running excel_shrink --only-clean-workbook on %s", file_path)
            subprocess.run(
                [
                    sys.executable,
                    str(PYTHON_EXCEL_SHRINK_SCRIPT),
                    str(file_path),
                    str(shrink_dir),
                    "--only-clean-workbook",
                ],
                check=True,
            )

            cleaned_path = shrink_dir / file_path.name
            if not cleaned_path.exists():
                raise FileNotFoundError(f"Cleaned workbook not found at {cleaned_path}")

            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.Application.DisplayAlerts = False

            wb = excel.Workbooks.Open(str(cleaned_path.resolve()))
            if wb is None:
                raise RuntimeError(f"Failed to open workbook: {cleaned_path}")

        def _save() -> None:
            nonlocal excel, wb
            assert wb is not None and excel is not None
            wb.SaveAs(str(out_path.resolve()))
            wb.Close()
            excel.Quit()

        return measure_open_save("Excel COM", _open, _save)


def _run_r_script(
    r_code: str, args: list[str]
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    if shutil.which("Rscript") is None:
        log.warning("Rscript not found – skipping R benchmark.")
        return None, None

    with tempfile.NamedTemporaryFile(
        delete=False, suffix=".R", mode="w", encoding="utf-8"
    ) as rf:
        rf.write(textwrap.dedent(r_code))
        r_path = Path(rf.name)

    try:
        completed = subprocess.run(
            ["Rscript", str(r_path), *args], capture_output=True, text=True, check=True
        )
        out = completed.stdout.strip()
        open_s, save_s = map(float, out.split(","))
        return open_s, save_s
    finally:
        try:
            r_path.unlink(missing_ok=True)
        except Exception:
            pass


def benchmark_r_openxlsx(
    file_path: Path, out_path: Path
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    r_code = """
    args <- commandArgs(trailingOnly=TRUE)
    infile  <- args[1]
    outfile <- args[2]
    library(openxlsx)

    t_start_open <- proc.time()
    wb <- tryCatch({ loadWorkbook(infile) }, error = function(e) { quit(status=1) })
    t_finish_open <- proc.time() - t_start_open
    open_time <- t_finish_open[["elapsed"]][[1]]

    t_start_save <- proc.time()
    tryCatch({ saveWorkbook(wb, outfile, overwrite = TRUE) }, error = function(e) { quit(status=1) })
    t_finish_save <- proc.time() - t_start_save
    save_time <- t_finish_save[["elapsed"]][[1]]

    cat(open_time, save_time, sep = ",")
    """
    return _run_r_script(r_code, [str(file_path), str(out_path)])


def benchmark_r_readxl_writexl(
    file_path: Path, out_path: Path
) -> tuple[Optional[Seconds], Optional[Seconds]]:
    r_code = """
    args <- commandArgs(trailingOnly = TRUE)
    file_in <- args[1]
    file_out <- args[2]

    library(readxl)
    library(writexl)

    sheets <- excel_sheets(file_in)

    t1 <- proc.time()
    data_list <- lapply(sheets, function(sht) read_excel(file_in, sheet = sht))
    t2 <- proc.time() - t1
    open_time <- t2["elapsed"][[1]]
    names(data_list) <- sheets

    t3 <- proc.time()
    write_xlsx(data_list, path=file_out)
    t4 <- proc.time() - t3
    save_time <- t4["elapsed"][[1]]

    cat(open_time, save_time, sep=",")
    """
    return _run_r_script(r_code, [str(file_path), str(out_path)])


@dataclass(frozen=True)
class ProcMeasure:
    stdout: Optional[str]
    stderr: Optional[str]
    duration_s: Seconds
    peak_private_bytes: Bytes
    returncode: int


def run_and_measure(
    cmd: list[str],
    *,
    capture_output: bool = False,
    text: bool = True,
    poll_interval: float = 0.1,
) -> ProcMeasure:
    if capture_output:
        proc = psutil.Popen(
            cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=text
        )
    else:
        proc = psutil.Popen(cmd)

    peak_private = 0
    t0 = time.perf_counter()

    def _safe_children(p: psutil.Process) -> list[psutil.Process]:
        try:
            return p.children(recursive=True)
        except (psutil.NoSuchProcess, psutil.AccessDenied):
            return []

    def _safe_private_bytes(p: psutil.Process) -> int:
        try:
            m = p.memory_full_info()
            return int(getattr(m, "private", 0) or 0)
        except (psutil.NoSuchProcess, psutil.AccessDenied, OSError):
            return 0

    stdout = stderr = None

    try:
        while True:
            total_private = _safe_private_bytes(proc)
            for child in _safe_children(proc):
                total_private += _safe_private_bytes(child)
            peak_private = max(peak_private, total_private)

            try:
                rc = proc.poll()
            except (psutil.NoSuchProcess, OSError):
                rc = 0
                break

            if rc is not None:
                break

            time.sleep(poll_interval)

        if capture_output:
            try:
                stdout, stderr = proc.communicate(timeout=5)
            except Exception:
                stdout, stderr = None, None
        else:
            try:
                proc.wait(timeout=5)
            except Exception:
                pass

        # one last peak update best-effort
        total_private = _safe_private_bytes(proc)
        for child in _safe_children(proc):
            total_private += _safe_private_bytes(child)
        peak_private = max(peak_private, total_private)

        duration = time.perf_counter() - t0

        try:
            returncode = proc.returncode
            if returncode is None:
                returncode = proc.wait(timeout=1)
        except (psutil.NoSuchProcess, OSError):
            returncode = 0

        return ProcMeasure(
            stdout=stdout,
            stderr=stderr,
            duration_s=duration,
            peak_private_bytes=peak_private,
            returncode=returncode,
        )

    except Exception:
        try:
            if proc.is_running():
                proc.kill()
        except (psutil.NoSuchProcess, psutil.AccessDenied, OSError):
            pass
        raise


BENCHMARKS: dict[
    str, Callable[[Path, Path], tuple[Optional[Seconds], Optional[Seconds]]]
] = {
    "openpyxl(default)": benchmark_openpyxl_default,
    "pandas": benchmark_pandas,
    "Microsoft Excel": benchmark_excel_com,
    "R openxlsx": benchmark_r_openxlsx,
    "R readxl+writexl": benchmark_r_readxl_writexl,
}


# ----------------------------
# Excel Shrink runners (controller)
# ----------------------------
@dataclass(frozen=True)
class ShrinkVariant:
    name: str
    extra_args: list[str]
    out_subdir: str


SHRINK_VARIANTS: list[ShrinkVariant] = [
    ShrinkVariant(DEFAULT_SHRINK_LIBRARY, [], "default"),
    ShrinkVariant(
        "Excel Shrink Rust (single-core)",
        EXCEL_SHRINK_SINGLE_CORE_ARGS,
        "single_core",
    ),
]


def run_excel_shrink_variant(
    original_file: Path,
    output_dir: Path,
    *,
    variant: ShrinkVariant,
    chunk_size: int = 2_097_152,
) -> tuple[Seconds, Path, Bytes, int]:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / original_file.name

    cmd = [
        sys.executable,
        str(EXCEL_SHRINK_SCRIPT),
        str(original_file),
        str(output_dir),
        "--chunk-size",
        str(chunk_size),
        "--force-overwrite",
    ]
    cmd.extend(EXCEL_SHRINK_RUST_DEFAULT_ARGS)
    cmd.extend(variant.extra_args)

    log.info("Running %s on %s", variant.name, original_file.name)
    m = run_and_measure(cmd, capture_output=False)

    if m.returncode == 0:
        log.info(
            "%s completed in %.2fs; peak private %.2f MiB",
            variant.name,
            m.duration_s,
            m.peak_private_bytes / (1024**2),
        )
    else:
        log.error(
            "%s failed (exit %s) for %s", variant.name, m.returncode, original_file.name
        )

    return m.duration_s, output_file, m.peak_private_bytes, m.returncode


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


# ----------------------------
# Controller
# ----------------------------
def controller_main(args: argparse.Namespace) -> None:
    script_path = Path(__file__).resolve()

    input_folder = Path(args.input_folder).resolve()
    shrunk_folder = Path(args.shrunk_folder).resolve()
    second_shrink_folder = Path(args.second_shrinkage_shrunk_folder).resolve()

    shrink_csv_out = Path(args.excel_shrink_benchmarks_csv).resolve()
    libs_csv_out = Path(args.excel_library_execution_times_csv).resolve()

    # Downloads (same behavior)
    files_to_download = [
        {
            "url": "https://www.destatis.de/DE/Themen/Staat/Oeffentliche-Finanzen/Ausgaben-Einnahmen/Publikationen/Downloads-Ausgaben-und-Einnahmen/statistischer-bericht-rechnungsergebnis-kernhaushalt-gemeinden-2140331217005.xlsx?__blob=publicationFile&v=4",
            "filename": "statistischer-bericht-kernhaushalt-gemeinden.xlsx",
        },
        {
            "url": "https://pasteur.epa.gov/uploads/10.23719/1503098/QT_RR%20data_dobutamine%20challenge%20test_Hazari_Dec%202015.xlsx",
            "filename": "QT_RR data_dobutamine challenge test_Hazari_Dec 2015.xlsx",
        },
        {
            "url": "https://www.destatis.de/DE/Themen/Gesellschaft-Umwelt/Bevoelkerung/Wanderungen/Publikationen/Downloads-Wanderungen/wanderungen-2010120217005.xlsx?__blob=publicationFile",
            "filename": "wanderungen-2010120217005.xlsx",
        },
    ]

    output_folder = SCRIPT_DIR / "input"
    output_folder.mkdir(parents=True, exist_ok=True)
    input_folder.mkdir(parents=True, exist_ok=True)

    # Ensure SSL certs
    import os

    os.environ["SSL_CERT_FILE"] = certifi.where()

    for item in files_to_download:
        out = output_folder / item["filename"]
        if out.exists():
            log.info("✅ File already exists: %s", out)
        else:
            log.info("⬇️ Downloading %s ...", item["filename"])
            try:
                urllib.request.urlretrieve(item["url"], str(out))
                log.info("✅ Downloaded: %s", out)
            except Exception as e:
                log.error("❌ Failed to download %s: %s", item["filename"], e)

    # Find input files
    excel_files = sorted(
        [
            p
            for p in input_folder.iterdir()
            if p.is_file() and p.suffix.lower() == ".xlsx"
        ]
    )
    log.info("Found %d Excel files in %s.", len(excel_files), input_folder)
    if not excel_files:
        log.warning("No .xlsx files found. Nothing to do.")
        return

    # Prepare output dirs
    shrunk_folder.mkdir(parents=True, exist_ok=True)
    second_shrink_folder.mkdir(parents=True, exist_ok=True)
    (SCRIPT_DIR / "shrunk_files_single_core").mkdir(parents=True, exist_ok=True)
    (SCRIPT_DIR / "shrunk_files_no_split").mkdir(parents=True, exist_ok=True)

    library_names = [
        "openpyxl(default)",
        "pandas",
        "Microsoft Excel",
        "R openxlsx",
        "R readxl+writexl",
    ]

    shrink_rows: list[dict[str, Any]] = []
    lib_rows: list[dict[str, Any]] = []

    # ✅ Single per-file loop (fixes the original bug)
    for original_path in excel_files:
        fname = original_path.name
        orig_size = maybe_size(original_path)

        # --- default shrink (used as “shrunk_path” baseline everywhere)
        default_variant = SHRINK_VARIANTS[0]
        default_out_dir = shrunk_folder
        shrink_time, shrunk_path, shrink_peak_mem, rc = run_excel_shrink_variant(
            original_path, default_out_dir, variant=default_variant
        )
        shrunk_size = maybe_size(shrunk_path) if (rc == 0) else None

        # second shrink on default-shrunk
        shrink_time_on_shrunk = None
        shrunk_size_after_second = None
        shrunk_shrink_peak_mem = None
        if rc == 0 and shrunk_path.exists():
            t2, shrunk2_path, pm2, rc2 = run_excel_shrink_variant(
                shrunk_path, second_shrink_folder, variant=default_variant
            )
            shrink_time_on_shrunk = t2 if rc2 == 0 else None
            shrunk_shrink_peak_mem = pm2 if rc2 == 0 else None
            shrunk_size_after_second = maybe_size(shrunk2_path) if rc2 == 0 else None

        shrink_rows.append(
            {
                "File": fname,
                "Program": default_variant.name,
                "Shrink Time (s)": shrink_time if rc == 0 else None,
                "Shrink Time on shrunk file (s)": shrink_time_on_shrunk,
                "Original Peak Memory": shrink_peak_mem if rc == 0 else None,
                "Shrinked Peak Memory": shrunk_shrink_peak_mem,
                "Original Size (bytes)": orig_size,
                "Shrinked Size (bytes)": shrunk_size,
                "Shrinked Size after Second Shrinkage (bytes)": shrunk_size_after_second,
            }
        )

        # --- other variants: run on original and on default-shrunk
        for variant in SHRINK_VARIANTS[1:]:
            if variant.out_subdir == "single_core":
                base_dir = SCRIPT_DIR / "shrunk_files_single_core"
            else:
                base_dir = SCRIPT_DIR / "shrunk_files_no_split"

            t_o, out_o, pm_o, rc_o = run_excel_shrink_variant(
                original_path, base_dir / "on_original", variant=variant
            )
            size_o = maybe_size(out_o) if rc_o == 0 else None

            t_s = None
            out_s = None
            pm_s = None
            size_s = None
            if rc == 0 and shrunk_path.exists():
                t_s2, out_s2, pm_s2, rc_s2 = run_excel_shrink_variant(
                    shrunk_path, base_dir / "on_default_shrunk", variant=variant
                )
                t_s = t_s2 if rc_s2 == 0 else None
                out_s = out_s2
                pm_s = pm_s2 if rc_s2 == 0 else None
                size_s = maybe_size(out_s2) if rc_s2 == 0 else None

            shrink_rows.append(
                {
                    "File": fname,
                    "Program": variant.name,
                    "Shrink Time (s)": t_o if rc_o == 0 else None,
                    "Shrink Time on shrunk file (s)": t_s,
                    "Original Peak Memory": pm_o if rc_o == 0 else None,
                    "Shrinked Peak Memory": pm_s,
                    "Original Size (bytes)": orig_size,
                    "Shrinked Size (bytes)": size_o,
                    "Shrinked Size after Second Shrinkage (bytes)": size_s,
                }
            )

        # --- measure libraries on original and default-shrunk
        for lib in library_names:
            orig_open, orig_save, orig_mem = call_measure_one(
                script_path, lib, original_path, on_shrunk=False
            )

            if rc == 0 and shrunk_path.exists():
                shr_open, shr_save, shr_mem = call_measure_one(
                    script_path, lib, shrunk_path, on_shrunk=True
                )
            else:
                shr_open = shr_save = shr_mem = None

            lib_rows.append(
                {
                    "File": fname,
                    "Library": lib,
                    "Original Open Time (s)": orig_open,
                    "Original Save Time (s)": orig_save,
                    "Original Peak Memory": orig_mem,
                    "Shrinked Open Time (s)": shr_open,
                    "Shrinked Save Time (s)": shr_save,
                    "Shrinked Peak Memory": shr_mem,
                }
            )

    shrink_fieldnames = [
        "File",
        "Program",
        "Shrink Time (s)",
        "Shrink Time on shrunk file (s)",
        "Original Peak Memory",
        "Shrinked Peak Memory",
        "Original Size (bytes)",
        "Shrinked Size (bytes)",
        "Shrinked Size after Second Shrinkage (bytes)",
    ]
    write_csv(shrink_csv_out, shrink_rows, shrink_fieldnames)
    log.info("Excel Shrink benchmarks saved to %s", shrink_csv_out)

    libs_fieldnames = [
        "File",
        "Library",
        "Original Open Time (s)",
        "Original Save Time (s)",
        "Original Peak Memory",
        "Shrinked Open Time (s)",
        "Shrinked Save Time (s)",
        "Shrinked Peak Memory",
    ]
    write_csv(libs_csv_out, lib_rows, libs_fieldnames)
    log.info("Excel library execution times saved to %s", libs_csv_out)

    generate_chart(str(libs_csv_out), str(shrink_csv_out))


def call_measure_one(
    script_path: Path, library_name: str, file_path: Path, *, on_shrunk: bool
) -> tuple[Optional[Seconds], Optional[Seconds], Optional[Bytes]]:
    cmd = [
        sys.executable,
        str(script_path),
        "measure-one",
        "--library",
        library_name,
        "--file",
        str(file_path),
    ]
    if on_shrunk:
        cmd.append("--on-shrunk")

    m = run_and_measure(cmd, capture_output=True, text=True)

    log.info(
        "%s completed in %.2fs; Peak memory %.2f MiB; file: %s",
        library_name,
        m.duration_s,
        m.peak_private_bytes / (1024**2),
        file_path.name,
    )

    if m.returncode != 0 or not m.stdout:
        log.error(
            "measure-one (%s) failed with exit %s; stderr: %r",
            library_name,
            m.returncode,
            m.stderr,
        )
        return None, None, None

    try:
        data = json.loads(m.stdout.strip())
        return (
            data.get("Original Open Time (s)"),
            data.get("Original Save Time (s)"),
            m.peak_private_bytes,
        )
    except json.JSONDecodeError:
        log.error("JSON parse error from measure-one output: %r", m.stdout)
        return None, None, None


# ----------------------------
# Charting (kept compatible with your CSV schema)
# ----------------------------
def generate_chart(
    excel_library_execution_times_csv: str, excel_shrink_benchmarks_csv: str
) -> None:
    log.info(
        "Generating chart from %s and %s",
        excel_library_execution_times_csv,
        excel_shrink_benchmarks_csv,
    )

    p_lib = Path(excel_library_execution_times_csv)
    p_shrink = Path(excel_shrink_benchmarks_csv)
    if not p_lib.exists():
        log.warning("Missing CSV: %s", p_lib)
        return
    if not p_shrink.exists():
        log.warning("Missing CSV: %s", p_shrink)
        return

    try:
        df_lib = pd.read_csv(p_lib)
        df_shrink = pd.read_csv(p_shrink)
        df_shrink = df_shrink.rename(columns={"Program": "Library"})
        _generate_complex_chart(df_lib=df_lib, df_shrink=df_shrink)
    except Exception:
        log.exception("Could not generate chart")


def _generate_complex_chart(df_lib: pd.DataFrame, df_shrink: pd.DataFrame) -> None:
    log.info("Using extended chart routine (with memory at bottom)...")

    if df_lib.empty and df_shrink.empty:
        log.warning("Both DataFrames are empty; nothing to plot.")
        return

    def _to_float(x: Any) -> float:
        try:
            return float(x) if pd.notna(x) else 0.0
        except Exception:
            return 0.0

    def _as_index(df: pd.DataFrame) -> Optional[pd.DataFrame]:
        if df.empty:
            return None
        return df.set_index(["File", "Library"], drop=False)

    def _value(idx: Optional[pd.DataFrame], f: str, lib: str, col: str) -> float:
        if idx is None or col not in idx.columns:
            return 0.0
        key = (f, lib)
        if key not in idx.index:
            return 0.0
        v = idx.loc[key, col]
        if isinstance(v, pd.Series):
            v = v.iloc[0]
        return _to_float(v)

    def _text_value(idx: Optional[pd.DataFrame], f: str, lib: str, col: str) -> str:
        if idx is None or col not in idx.columns:
            return ""
        key = (f, lib)
        if key not in idx.index:
            return ""
        v = idx.loc[key, col]
        if isinstance(v, pd.Series):
            v = v.iloc[0]
        if pd.isna(v):
            return ""
        return str(v)

    files = sorted(
        set(df_lib.get("File", pd.Series([], dtype=str)).dropna().unique()).union(
            set(df_shrink.get("File", pd.Series([], dtype=str)).dropna().unique())
        )
    )
    if not files:
        log.info("No files found in data; nothing to plot.")
        return

    file_labels: dict[str, str] = {f: f"File {i+1}" for i, f in enumerate(files)}
    n_files = len(files)

    # --- layout / label metrics
    PX_PER_CHAR = 6.0
    INSIDE_PAD_PX = 0.0
    RIGHT_LABEL_PAD_PX = 12.0
    RIGHT_ARROW_SHRINK = 1
    FONT_SIZE = 7
    VERTICAL_LABEL_OFFSET_PX = -0.6

    def _data_per_px_x(ax: plt.Axes) -> float:
        x_min, x_max = ax.get_xlim()
        return (x_max - x_min) / max(ax.bbox.width, 1.0)

    def _data_per_px_y(ax: plt.Axes) -> float:
        y_min, y_max = ax.get_ylim()
        return (y_max - y_min) / max(ax.bbox.height, 1.0)

    def can_fit_inside(ax: plt.Axes, width_data: float, text: str) -> bool:
        needed_data = ((len(text) * PX_PER_CHAR) + INSIDE_PAD_PX) * _data_per_px_x(ax)
        return width_data >= needed_data

    def right_label(ax: plt.Axes, x_end: float, y: float, parts: list[str]) -> None:
        x_text = x_end + RIGHT_LABEL_PAD_PX * _data_per_px_x(ax)
        text = ", ".join(parts)
        ax.annotate(
            text,
            xy=(x_end, y + 0.001),
            xytext=(x_text, y + 0.008),
            va="center",
            ha="left",
            fontsize=FONT_SIZE,
            arrowprops=dict(arrowstyle="->", shrinkA=RIGHT_ARROW_SHRINK, lw=0.8),
        )

    def text_vshift(ax: plt.Axes, points: float) -> mtransforms.ScaledTranslation:
        return mtransforms.ScaledTranslation(
            0, points / 72.0, ax.figure.dpi_scale_trans
        )

    @dataclass(frozen=True)
    class Segment:
        width: float
        color: str
        label: str  # legend label (can be "")
        hatch: str = ""

    def draw_stacked_barh(
        ax: plt.Axes,
        *,
        y: float,
        segments: list[Segment],
        height: float,
        start: float = 0.0,
        inside_labels: Optional[list[str]] = None,
    ) -> None:
        """
        Draw stacked horizontal bars and place labels either inside when they fit,
        otherwise as a single right-side callout.
        """
        x = start
        inside_labels = inside_labels or ["" for _ in segments]
        assert len(inside_labels) == len(segments)

        labels_ok: list[bool] = []
        for seg, seg_label in zip(segments, inside_labels):
            if seg.width <= 0:
                labels_ok.append(False)
                continue
            ax.barh(
                y,
                seg.width,
                height=height,
                left=x,
                color=seg.color,
                label=seg.label,
                hatch=seg.hatch,
                edgecolor="black" if seg.hatch else None,
            )
            ok = bool(seg_label) and can_fit_inside(ax, seg.width, seg_label)
            labels_ok.append(ok)
            if ok:
                ax.text(
                    x + seg.width / 2.0,
                    y,
                    seg_label,
                    va="center",
                    ha="center",
                    fontsize=FONT_SIZE,
                    color="white",
                    transform=ax.transData + text_vshift(ax, VERTICAL_LABEL_OFFSET_PX),
                )
            x += seg.width

        # If any label didn't fit, render a right-side label with non-zero parts.
        if not all(labels_ok):
            parts = [
                lab
                for seg, lab in zip(segments, inside_labels)
                if seg.width > 0 and lab
            ]
            if parts:
                right_label(ax, start + sum(s.width for s in segments), y, parts)

    def draw_failure_barh(
        ax: plt.Axes,
        *,
        y: float,
        duration: float,
        height: float,
        label: str,
        legend_label: str = "",
        start: float = 0.0,
    ) -> None:
        if duration <= 0:
            duration = 0.01
        ax.barh(
            y,
            duration,
            height=height,
            left=start,
            color="#D62728",
            edgecolor="black",
            hatch="////",
            label=legend_label,
            alpha=0.92,
        )
        right_label(ax, start + duration, y, [label])

    def _can_fit_inside_vertical(
        ax: plt.Axes, bar_height_data: float, font_size_pts: int = 7, fudge: float = 1.1
    ) -> bool:
        text_px = (font_size_pts / 72.0) * ax.figure.dpi * fudge
        needed_data = text_px * _data_per_px_y(ax)
        return bar_height_data >= needed_data

    lib_idx = _as_index(df_lib)
    shr_idx = _as_index(df_shrink)

    def lib_value(f: str, lib: str, col: str) -> float:
        return _value(lib_idx, f, lib, col)

    def lib_text(f: str, lib: str, col: str) -> str:
        return _text_value(lib_idx, f, lib, col)

    def shrink_value(f: str, lib: str, col: str) -> float:
        return _value(shr_idx, f, lib, col)

    # --- series definitions
    library_info: list[tuple[str, str, str]] = [
        ("openpyxl(default)", "darkorange", "orangered"),
        ("pandas", "steelblue", "royalblue"),
        ("Microsoft Excel", "darkgreen", "limegreen"),
        ("R openxlsx", "red", "firebrick"),
        ("R readxl+writexl", "chocolate", "sienna"),
        (DEFAULT_SHRINK_LIBRARY, "mediumpurple", "purple"),
    ]
    n_libs = len(library_info)

    extra_run_info: list[tuple[str, str]] = [
        ("Excel Shrink Rust (single-core)", "slateblue"),
    ]
    n_extra = len(extra_run_info)

    fig = plt.figure(figsize=(14, 7 + 1.8 * n_files))
    gs = GridSpec(nrows=2, ncols=1, height_ratios=[7, 1.8 * n_files], hspace=0.20)

    gs_top = GridSpecFromSubplotSpec(
        nrows=2, ncols=1, subplot_spec=gs[0], height_ratios=[4, 1], hspace=0.35
    )
    ax_time = fig.add_subplot(gs_top[0, 0])
    ax_size = fig.add_subplot(gs_top[1, 0])

    gs_bottom = GridSpecFromSubplotSpec(
        nrows=n_files, ncols=1, subplot_spec=gs[1], hspace=0.08
    )
    mem_axes = [fig.add_subplot(gs_bottom[i, 0]) for i in range(n_files)]
    mem_axes[0].set_title("Memory Consumption")
    plt.subplots_adjust(left=0.18)

    total_bars_per_file = (n_libs - 1) + n_libs + n_extra
    bar_height = 0.09
    offsets_all = np.linspace(-0.60, 0.60, total_bars_per_file)
    offsets_orig = offsets_all[: (n_libs - 1)]
    offsets_shr = offsets_all[(n_libs - 1) : (n_libs - 1 + n_libs)]
    offsets_extra = offsets_all[(n_libs - 1 + n_libs) :]

    y_positions_top = np.arange(n_files) * (bar_height * (total_bars_per_file + 2))

    # Precompute once (used repeatedly)
    libraries_no_shrink = [
        (n, co, cs) for (n, co, cs) in library_info if n != DEFAULT_SHRINK_LIBRARY
    ]

    # -----------------
    # TIME section
    # -----------------
    for i, f in enumerate(files):
        base = float(y_positions_top[i])
        shrink_t = shrink_value(f, DEFAULT_SHRINK_LIBRARY, "Shrink Time (s)")

        # Original runs (libraries excluding Excel Shrink)
        for lib_index, (lib_name, c_open, c_save) in enumerate(libraries_no_shrink):
            pos = base + float(offsets_orig[lib_index])
            open_t = lib_value(f, lib_name, "Original Open Time (s)")
            save_t = lib_value(f, lib_name, "Original Save Time (s)")
            original_status = lib_text(f, lib_name, "Original Status")
            original_failure_t = lib_value(f, lib_name, "Original Failure Time (s)")

            if original_status and original_status != "ok":
                draw_failure_barh(
                    ax_time,
                    y=pos,
                    height=bar_height,
                    duration=original_failure_t,
                    label=f"{lib_name} failed/OOM after {original_failure_t:.1f}s",
                    legend_label="Failed / OOM",
                )
                continue

            draw_stacked_barh(
                ax_time,
                y=pos,
                height=bar_height,
                segments=[
                    Segment(open_t, c_open, f"{lib_name} open" if i == 0 else ""),
                    Segment(save_t, c_save, f"{lib_name} save" if i == 0 else ""),
                ],
                inside_labels=[
                    f"{open_t:.1f}" if open_t > 0 else "",
                    f"{save_t:.1f}" if save_t > 0 else "",
                ],
            )

        # Shrunk runs: prepend shrink time, then either (a) second shrink run or (b) open+save on shrunk
        for lib_index, (lib_name, c_open, c_save) in enumerate(library_info):
            pos = base + float(offsets_shr[lib_index])

            # Always start with shrink time segment
            if lib_index == 0 and i == 0:
                shrink_label = "Shrink time"
            else:
                shrink_label = ""

            if lib_name == DEFAULT_SHRINK_LIBRARY:
                shrink_t2 = shrink_value(
                    f, DEFAULT_SHRINK_LIBRARY, "Shrink Time on shrunk file (s)"
                )
                draw_stacked_barh(
                    ax_time,
                    y=pos,
                    height=bar_height,
                    segments=[
                        Segment(shrink_t, "mediumorchid", shrink_label),
                        Segment(
                            shrink_t2,
                            c_open,
                            f"{lib_name} shrunk open" if i == 0 else "",
                        ),
                    ],
                    inside_labels=[
                        f"{shrink_t:.1f}" if shrink_t > 0 else "",
                        f"{shrink_t2:.1f}" if shrink_t2 > 0 else "",
                    ],
                )
            else:
                shr_open = lib_value(f, lib_name, "Shrinked Open Time (s)")
                shr_save = lib_value(f, lib_name, "Shrinked Save Time (s)")
                shr_status = lib_text(f, lib_name, "Shrinked Status")
                shr_failure_t = lib_value(f, lib_name, "Shrinked Failure Time (s)")

                if shr_status and shr_status != "ok":
                    draw_stacked_barh(
                        ax_time,
                        y=pos,
                        height=bar_height,
                        segments=[Segment(shrink_t, "mediumorchid", shrink_label)],
                        inside_labels=[f"{shrink_t:.1f}" if shrink_t > 0 else ""],
                    )
                    draw_failure_barh(
                        ax_time,
                        y=pos,
                        height=bar_height,
                        duration=shr_failure_t,
                        start=shrink_t,
                        label=f"{lib_name} failed/OOM after {shr_failure_t:.1f}s",
                        legend_label="Failed / OOM",
                    )
                    continue

                draw_stacked_barh(
                    ax_time,
                    y=pos,
                    height=bar_height,
                    segments=[
                        Segment(shrink_t, "mediumorchid", shrink_label),
                        Segment(
                            shr_open,
                            c_open,
                            f"{lib_name} shrunk open" if i == 0 else "",
                        ),
                        Segment(
                            shr_save,
                            c_save,
                            f"{lib_name} shrunk save" if i == 0 else "",
                        ),
                    ],
                    inside_labels=[
                        f"{shrink_t:.1f}" if shrink_t > 0 else "",
                        f"{shr_open:.1f}" if shr_open > 0 else "",
                        f"{shr_save:.1f}" if shr_save > 0 else "",
                    ],
                )

        # Extra shrink variants (single segment)
        for ex_i, (run_name, run_color) in enumerate(extra_run_info):
            pos = base + float(offsets_extra[ex_i])
            run_t = shrink_value(f, run_name, "Shrink Time (s)")
            draw_stacked_barh(
                ax_time,
                y=pos,
                height=bar_height,
                segments=[Segment(run_t, run_color, run_name if i == 0 else "")],
                inside_labels=[f"{run_t:.1f}" if run_t > 0 else ""],
            )

    ax_time.margins(x=0.02)
    ax_time.set_yticks(y_positions_top)
    ax_time.set_yticklabels([file_labels[f] for f in files])
    ax_time.invert_yaxis()
    ax_time.set_xlabel("Time (seconds)")
    ax_time.set_title("Time Comparison per Excel File")

    handles, labels = ax_time.get_legend_handles_labels()
    pairs = [(h, l) for h, l in zip(handles, labels) if l]
    seen: "OrderedDict[str, Any]" = OrderedDict()
    for h, l in pairs:
        if DEFAULT_SHRINK_LIBRARY in l and "shrunk" in l:
            l = l.replace(
                f"{DEFAULT_SHRINK_LIBRARY} shrunk save",
                "Second Excel Shrink Rust run",
            )
        if "shrunk" in l:
            continue
        if l not in seen:
            seen[l] = h
    ax_time.legend(list(seen.values()), list(seen.keys()), ncol=3, fontsize=8)

    # -----------------
    # SIZE section
    # -----------------
    default_sizes = df_shrink[df_shrink["Library"] == DEFAULT_SHRINK_LIBRARY].copy()
    size_by_file = default_sizes.set_index("File")[
        ["Original Size (bytes)", "Shrinked Size (bytes)"]
    ]

    def _bytes_for_file(f: str) -> tuple[float, float]:
        if f not in size_by_file.index:
            return 0.0, 0.0
        row = size_by_file.loc[f]
        if isinstance(row, pd.DataFrame):
            row = row.iloc[0]
        return _to_float(row.get("Original Size (bytes)", 0)), _to_float(
            row.get("Shrinked Size (bytes)", 0)
        )

    orig_size_b_list: list[float] = []
    shrunk_size_b_list: list[float] = []
    orig_size_mb_list: list[float] = []
    shrunk_size_mb_list: list[float] = []

    for f in files:
        orig_b, shr_b = _bytes_for_file(f)
        orig_size_b_list.append(orig_b)
        shrunk_size_b_list.append(shr_b)
        orig_size_mb_list.append(orig_b / (1024 * 1024))
        shrunk_size_mb_list.append(shr_b / (1024 * 1024))

    def _bar_label(ax: plt.Axes, x: float, y: float, text: str, fontsize: int) -> None:
        ax.text(x, y, text, va="center", ha="center", fontsize=fontsize, color="white")

    y_positions_bottom = np.arange(n_files) * 0.1
    file_bar_height = 0.04

    for i, f in enumerate(files):
        base = float(y_positions_bottom[i])
        ax_size.barh(
            base - file_bar_height * 0.55,
            orig_size_mb_list[i],
            height=file_bar_height,
            color="darkorange",
            label="Original Size (MB)" if i == 0 else "",
        )
        ax_size.barh(
            base + file_bar_height * 0.55,
            shrunk_size_mb_list[i],
            height=file_bar_height,
            color="steelblue",
            label="Shrinked Size (MB)" if i == 0 else "",
        )

        _bar_label(
            ax_size,
            max(orig_size_mb_list[i] * 0.97, 0.01),
            base - file_bar_height * 0.55,
            f"{orig_size_mb_list[i]:.1f} MB",
            FONT_SIZE,
        )

    for i, f in enumerate(files):
        base = float(y_positions_bottom[i])
        orig_b = orig_size_b_list[i]
        shr_b = shrunk_size_b_list[i]
        shr_mb = shrunk_size_mb_list[i]

        if orig_b > 0:
            red_pct = 100.0 * (1.0 - (shr_b / orig_b))
            label = f"{shr_mb:.1f} MB (−{red_pct:.1f}%)"

            x_text = max(
                shr_mb + 10 * _data_per_px_x(ax_size),
                ax_size.get_xlim()[0] + 5 * _data_per_px_x(ax_size),
            )

            ax_size.annotate(
                label,
                xy=(shr_mb + 0.15, base + file_bar_height * 0.55),
                xytext=(x_text, base + file_bar_height * 0.55),
                va="center",
                ha="left",
                fontsize=FONT_SIZE,
                arrowprops=dict(arrowstyle="->", lw=0.8, shrinkA=0, shrinkB=0),
            )

    ax_size.set_yticks(y_positions_bottom)
    ax_size.set_yticklabels([file_labels[f] for f in files])
    ax_size.invert_yaxis()
    ax_size.set_xlabel("File Size (MB)")
    ax_size.set_title("File Size Comparison per Excel File")
    ax_size.legend(ncol=2, fontsize=8)

    # -----------------
    # MEMORY section
    # -----------------
    mem_libraries = [
        "openpyxl(default)",
        "pandas",
        "Microsoft Excel",
        "R openxlsx",
        "R readxl+writexl",
        DEFAULT_SHRINK_LIBRARY,
        "Excel Shrink Rust (single-core)",
    ]

    if not df_lib.empty:
        df_lib["Original Peak Memory"] = pd.to_numeric(
            df_lib.get("Original Peak Memory", 0), errors="coerce"
        ).fillna(0.0)
        df_lib["Shrinked Peak Memory"] = pd.to_numeric(
            df_lib.get("Shrinked Peak Memory", 0), errors="coerce"
        ).fillna(0.0)
    if not df_shrink.empty:
        df_shrink["Original Peak Memory"] = pd.to_numeric(
            df_shrink.get("Original Peak Memory", 0), errors="coerce"
        ).fillna(0.0)
        df_shrink["Shrinked Peak Memory"] = pd.to_numeric(
            df_shrink.get("Shrinked Peak Memory", 0), errors="coerce"
        ).fillna(0.0)

    def mem_pair_mib(file_name: str, lib: str) -> tuple[float, float]:
        if lib.startswith("Excel Shrink"):
            o = shrink_value(file_name, lib, "Original Peak Memory")
            s = shrink_value(file_name, lib, "Shrinked Peak Memory")
        else:
            o = lib_value(file_name, lib, "Original Peak Memory")
            s = lib_value(file_name, lib, "Shrinked Peak Memory")
        return o / (1024**2), s / (1024**2)

    def pretty_mem_label(lib: str) -> str:
        return lib.replace(
            "Excel Shrink Rust (single-core)", "Excel Shrink Rust\n(single-core)"
        )

    for ax, file_name in zip(mem_axes, files):
        x = np.arange(len(mem_libraries))
        bar_width = 0.35

        orig_mem_mib: list[float] = []
        shr_mem_mib: list[float] = []
        for lib in mem_libraries:
            o, s = mem_pair_mib(file_name, lib)
            orig_mem_mib.append(o)
            shr_mem_mib.append(s)

        bars_orig = ax.bar(
            x - bar_width / 2,
            orig_mem_mib,
            width=bar_width,
            label="Original Peak Memory (MiB)",
            edgecolor="black",
        )
        bars_shr = ax.bar(
            x + bar_width / 2,
            shr_mem_mib,
            width=bar_width,
            label="Shrinked Peak Memory (MiB)",
            edgecolor="black",
        )

        PAD_PX = 6
        y_pad = PAD_PX * _data_per_px_y(ax)

        for bar in list(bars_orig) + list(bars_shr):
            h = float(bar.get_height())
            if h <= 0:
                continue
            label = f"{int(round(h))}"
            x_center = float(bar.get_x() + bar.get_width() / 2.0)
            if _can_fit_inside_vertical(ax, h, font_size_pts=FONT_SIZE):
                ax.text(
                    x_center,
                    h / 2.0,
                    label,
                    ha="center",
                    va="center",
                    fontsize=FONT_SIZE,
                    color="white",
                )
            else:
                ax.annotate(
                    label,
                    xy=(x_center, h),
                    xytext=(x_center, h + y_pad),
                    ha="center",
                    va="bottom",
                    fontsize=FONT_SIZE,
                    arrowprops=dict(arrowstyle="->", lw=0.8, shrinkA=0, shrinkB=0),
                )

        ax.set_ylabel(f"Peak Memory (MiB)\n{file_labels[file_name]}", fontsize=9)
        ax.ticklabel_format(style="plain", axis="y")
        ax.set_xticks(x)
        ax.set_xticklabels(
            [pretty_mem_label(lib) for lib in mem_libraries], ha="center", fontsize=8
        )
        ax.legend(fontsize=8)
        ax.grid(axis="x", linestyle="--", alpha=0.5)

    plt.tight_layout()
    CHART_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(CHART_OUTPUT_PATH, dpi=300, bbox_inches="tight")
    log.info("Chart generated and saved to %s", CHART_OUTPUT_PATH)
    plt.close(fig)


# ----------------------------
# CLI
# ----------------------------
def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run measurements on Excel files and generate comparison chart."
    )
    subparsers = parser.add_subparsers(
        dest="mode", required=True, help="Subcommand: controller or measure-one"
    )

    ctrl_parser = subparsers.add_parser(
        "controller", help="Run measurements on all files and libraries."
    )
    ctrl_parser.add_argument(
        "--run-tests",
        "--run_tests",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Whether to run the measurements (default: True).",
    )

    meas_parser = subparsers.add_parser(
        "measure-one", help="Measure one library on one file."
    )
    meas_parser.add_argument(
        "--library", required=True, help="Library name, e.g. 'openpyxl(default)'"
    )
    meas_parser.add_argument("--file", required=True, help="Path to the .xlsx file")
    meas_parser.add_argument(
        "--on-shrunk",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Whether the file is a shrunk file (default: False).",
    )

    args, _unknown = parser.parse_known_args()

    # keep same default folder layout as your script
    input_folder = SCRIPT_DIR / "input"
    shrunk_folder = SCRIPT_DIR / "shrunk_files"
    second_shrink_folder = SCRIPT_DIR / "second_shrinkage_shrunk_files"
    shrink_csv = SCRIPT_DIR / "excel_shrink_benchmarks.csv"
    libs_csv = SCRIPT_DIR / "excel_library_execution_times.csv"

    if args.mode == "measure-one":
        measure_one_main(args)
        return

    # controller
    args.input_folder = str(input_folder)
    args.shrunk_folder = str(shrunk_folder)
    args.second_shrinkage_shrunk_folder = str(second_shrink_folder)
    args.excel_shrink_benchmarks_csv = str(shrink_csv)
    args.excel_library_execution_times_csv = str(libs_csv)

    if (
        not args.run_tests
        and Path(args.excel_shrink_benchmarks_csv).exists()
        and Path(args.excel_library_execution_times_csv).exists()
    ):
        log.info(
            "CSV files already exist. Skipping measurements and only generating chart."
        )
        generate_chart(
            args.excel_library_execution_times_csv, args.excel_shrink_benchmarks_csv
        )
        sys.exit(0)

    controller_main(args)


if __name__ == "__main__":
    main()
